"""Parse PDF, DOCX and PPTX bytes into text sections that keep page/slide numbers."""

import io
import re
from dataclasses import dataclass
from pathlib import PurePosixPath

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx"}


@dataclass(frozen=True)
class Section:
    text: str
    page: int | None  # 1-based PDF page or PPTX slide; None for DOCX


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def load_document(filename: str, data: bytes) -> list[Section]:
    ext = PurePosixPath(filename).suffix.lower()
    if ext == ".pdf":
        sections = _load_pdf(data)
    elif ext == ".docx":
        sections = _load_docx(data)
    elif ext == ".pptx":
        sections = _load_pptx(data)
    else:
        raise ValueError(f"Unsupported file type: {filename}")
    return [s for s in sections if s.text]


def _load_pdf(data: bytes) -> list[Section]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return [Section(_clean(page.extract_text() or ""), i) for i, page in enumerate(reader.pages, 1)]


def _load_docx(data: bytes) -> list[Section]:
    from docx import Document
    from docx.table import Table

    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    for block in doc.iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
            parts.append("")
            continue
        text = block.text.strip()
        if not text:
            continue
        style = (block.style.name if block.style is not None else "") or ""
        if style.startswith("Heading") or style == "Title":
            parts.append(f"\n## {text}")
        else:
            parts.append(text)
    return [Section(_clean("\n".join(parts)), None)]


def _pptx_shape_texts(shapes) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    texts: list[str] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            texts.extend(_pptx_shape_texts(shape.shapes))
        elif shape.has_text_frame and shape.text_frame.text.strip():
            texts.append(shape.text_frame.text.strip())
        elif getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    texts.append(" | ".join(cells))
    return texts


def _load_pptx(data: bytes) -> list[Section]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    sections = []
    for i, slide in enumerate(prs.slides, 1):
        texts = _pptx_shape_texts(slide.shapes)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                texts.append(f"Speaker notes: {notes}")
        sections.append(Section(_clean("\n".join(texts)), i))
    return sections
